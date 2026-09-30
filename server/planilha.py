"""Exportação em planilha no padrão da exportação do Coletum (xlsx e csv).

Reproduz as regras da exportação do Coletum (xlsx e csv). Única divergência de propósito: coordenada com até
5 casas decimais, como a tela do Coletum (o serviço de exportação escreve 12).

Fluxo: montar_esquema (estrutura + ajustes, antes de ler a API) -> preencher (preenchimentos) -> gravar_xlsx ou
gravar_csv. Os ajustes pedidos na conversa (campos, aba única, metadados, rótulos, CSV, extras) mudam o
esquema; sem ajustes sai o padrão do Coletum.
"""
from __future__ import annotations

import re
import unicodedata
from dataclasses import dataclass, field
from datetime import datetime, timezone
from decimal import Decimal
from pathlib import Path
from zoneinfo import ZoneInfo

from achatar import coord5

MIDIA = {"gallery", "photo", "file"}
SEMPRE_MULTI = MIDIA | {"checkbox"}          # sempre viram aba própria
NUNCA_MULTI = {"rating", "signature", "agreement"}
CSV_EM_LINHA = {"select", "radio", "rating", "signature", "agreement"}  # no CSV, ficam na linha
COORD = {"coordinate", "geolocation"}
FONTES = {"mobile": "Aplicativo móvel", "web_private": "Sistema web", "web_public": "Link público"}
SIMBOLO_MOEDA = {"BRL": "R$", "USD": "US$", "EUR": "€", "GBP": "£"}
VERDE, BRANCO = "FF1CAF9A", "FFFFFFFF"
GERADO_POR_PADRAO = "Coletum via MCP"
TOPO_LEIA_ME = 26

MODOS_ABA_UNICA = {"linhas", "colunas"}
POSICOES_META = {"fim", "inicio", "fora"}
EXTRAS = {"precisao": "Local da coleta (Precisão em metros)", "altitude": "Local da coleta (Altitude em metros)",
          "origem": "Origem"}
SINONIMOS_EXTRAS = {"precisão": "precisao", "precision": "precisao", "elevation": "altitude", "fonte": "origem",
                    "plataforma": "origem"}
CHAVES_AJUSTES = {"campos", "aba_unica", "metadados", "rotulos", "csv", "extras"}


class ErroAjuste(ValueError):
    """Ajuste inválido, recusado antes de ler os preenchimentos."""


# --------------------------------------------------------------------------------------------
# Texto: nomes de aba e de arquivo
# --------------------------------------------------------------------------------------------

def sem_acento(s: str) -> str:
    d = unicodedata.normalize("NFD", s)
    return unicodedata.normalize("NFC", "".join(ch for ch in d if unicodedata.category(ch) != "Mn"))


def nome_aba(s: str) -> str:
    s = sem_acento(s or "")
    for ch in (":", "\\", "/", "?", "*", "[", "]", "'"):
        s = s.replace(ch, "")
    nome = s.strip()[:30].rstrip(" ")
    return nome or "Sheet"


def nomes_unicos(abas: list) -> None:
    usados: set[str] = set()
    for a in abas:
        chave = a.nome.strip().lower()
        if chave not in usados:
            usados.add(chave)
            continue
        base, n = a.nome[:27], 1
        while True:
            cand = f"{base}_{n:02d}"
            if cand.strip().lower() not in usados:
                a.nome = cand
                usados.add(cand.strip().lower())
                break
            n += 1


def slugify(s: str) -> str:
    s = sem_acento(s).lower()
    saida, hifen = [], True
    for ch in s:
        if "a" <= ch <= "z" or "0" <= ch <= "9":
            saida.append(ch)
            hifen = False
        elif not hifen:
            saida.append("-")
            hifen = True
    return "".join(saida).rstrip("-")


def nome_arquivo(nome: str, versao: str, quando: datetime) -> str:
    return slugify(f"{(nome or '')[:30]}_v{versao or ''}_{quando.strftime('%d-%m-%Y-%H-%M-%S')}")


# --------------------------------------------------------------------------------------------
# Valores
# --------------------------------------------------------------------------------------------

@dataclass
class Cel:
    """Célula tipada; o xlsx e o csv formatam cada tipo do seu jeito."""
    tipo: str                 # texto | num | data | datahora
    valor: object = None
    xfmt: str | None = None   # máscara do Excel (inteiro, decimal, moeda)
    csv: str = "dec"          # num no csv: dec (separador decimal escolhido) | int (sem troca) | moeda
    moeda: str = ""


VAZIA = None


def texto(v) -> Cel | None:
    s = str(v).replace("\r\n", "\n").replace("\r", "\n")
    return Cel("texto", s) if s != "" else VAZIA


def go_f(v: float) -> str:
    """strconv.FormatFloat(v, 'f', -1, 64): menor representação exata, sem expoente."""
    if v == int(v) and abs(v) < 1e16:
        return str(int(v))
    s = format(Decimal(repr(float(v))), "f")
    return s.rstrip("0").rstrip(".") if "." in s else s


def go_v(v) -> str:
    """fmt.Sprint de um valor vindo do JSON."""
    if isinstance(v, bool):
        return "true" if v else "false"
    if isinstance(v, float):
        return go_f(v)
    if isinstance(v, list):
        return "; ".join(go_v(x) for x in v)
    return str(v)


def num(v) -> float | None:
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    return None


def arredonda_coord(v: float) -> float:
    """Até 5 casas, como a tela do Coletum (o serviço de exportação escreve 12).
    A célula do xlsx continua numérica, no formato Geral, que não mostra zeros à direita."""
    r = coord5(v)
    return 0.0 if r is None else r


def mascara_decimal(v: float) -> str:
    s = go_f(v)
    casas = max(2, len(s) - s.index(".") - 1) if "." in s else 2
    return "#,##0." + "0" * casas


def ler_data(v) -> datetime | None:
    """Datas do Webservice V2: com fuso (com ou sem dois-pontos, ou Z), sem fuso (lida como UTC) ou só o dia."""
    if not isinstance(v, str) or not v:
        return None
    s = v.strip()
    try:
        dt = datetime.fromisoformat(s.replace("Z", "+00:00"))
    except ValueError:
        return None
    return dt if dt.tzinfo is not None else dt.replace(tzinfo=timezone.utc)


def ler_dia(v) -> datetime | None:
    """Campo data: o dia como veio (relógio de parede), sem converter fuso."""
    if not isinstance(v, str) or not v:
        return None
    try:
        dt = datetime.fromisoformat(v.strip().replace("Z", "+00:00"))
    except ValueError:
        return None
    return datetime(dt.year, dt.month, dt.day)


def celula_folha(tipo: str, v, fuso: ZoneInfo) -> Cel | None:
    if v is None:
        return VAZIA
    if tipo == "int":
        n = num(v) or 0.0
        return Cel("num", n, "0" if n else None, "int")
    if tipo in ("float", "number"):
        n = num(v) or 0.0
        return Cel("num", n, mascara_decimal(n) if n else None, "dec")
    if tipo in ("range", "rating"):
        return Cel("num", num(v) or 0.0, None, "dec" if tipo == "range" else "int")
    if tipo == "date":
        d = ler_dia(v)
        return Cel("data", d) if d else texto(go_v(v))
    if tipo == "datetime":
        d = ler_data(v)
        return Cel("datahora", d.astimezone(fuso)) if d else texto(go_v(v))
    if tipo == "agreement":
        return Cel("texto", "Sim" if v else "Não") if isinstance(v, bool) else VAZIA
    if tipo == "money":
        if not isinstance(v, dict):
            return VAZIA
        cod = v.get("currency") if isinstance(v.get("currency"), str) else ""
        return Cel("num", num(v.get("value")) or 0.0, None, "moeda", cod)
    if tipo == "relational":
        return texto(v.get("label") or "") if isinstance(v, dict) else VAZIA
    if tipo == "signature":
        return texto(v) if isinstance(v, str) else VAZIA
    return texto(go_v(v))


def simbolo(moeda: str) -> str | None:
    if moeda in SIMBOLO_MOEDA:
        return SIMBOLO_MOEDA[moeda]
    return moeda if re.fullmatch(r"[A-Z]{3}", moeda or "") else None


def partes_midia(url: str) -> tuple[str, str]:
    """Nome original (fim do link) e nome do arquivo (nome-pasta.ext), como a exportação."""
    p = re.split(r"[?#]", url, 1)[0].rstrip("/")
    segs = p.split("/")
    base = segs[-1]
    pasta = segs[-2] if len(segs) >= 2 else ""
    nome, ext = (base.rsplit(".", 1) + [""])[:2] if "." in base else (base, "")
    return base, (f"{nome}-{pasta}.{ext}" if ext else f"{nome}-{pasta}")


def coord(v) -> tuple[float, float] | None:
    if isinstance(v, dict):
        c = v.get("coordinates")
        if isinstance(c, (list, tuple)) and len(c) >= 2:
            return float(num(c[1]) or 0.0), float(num(c[0]) or 0.0)
    return None


def relacional(v) -> tuple[str, str] | None:
    if isinstance(v, dict):
        rid = v.get("friendlyId") or v.get("answer_id") or ""
        return str(rid), str(v.get("label") or "")
    return None


def rel_tem_rotulo(c: dict) -> bool:
    """O Webservice V2 manda as opções do relacional como texto 'código - rótulo' (ou só o código, sem campos
    de exibição); a exportação usa o rótulo da opção para decidir se existe a coluna (Rótulo)."""
    for o in c.get("options") or []:
        if isinstance(o, dict) and o.get("label"):
            return True
        if isinstance(o, str) and " - " in o and o.split(" - ", 1)[1].strip():
            return True
    return False


# --------------------------------------------------------------------------------------------
# Esquema
# --------------------------------------------------------------------------------------------

@dataclass(eq=False)
class Col:
    rotulo: str
    tipo: str                      # cod_pai cod direto lat lon qtd ass_nome ass_link ass_arq rel_cod rel_rot
                                   # mid_nome mid_link mid_arq item meta
    comp: dict | None = None
    caminho: tuple = ()            # ids de grupos não repetíveis entre o escopo da linha e o campo
    filha: "Aba | None" = None
    anc: tuple = ()                # componentes do grupo mais externo até o próprio campo
    meta: str = ""


@dataclass(eq=False)
class Aba:
    nome: str
    rotulo: str
    cols: list = field(default_factory=list)
    comp: dict | None = None
    grupo: bool = False
    pai: "Aba | None" = None
    caminho_topo: tuple = ()       # para abas ligadas à raiz
    caminho_aninhado: tuple = ()   # do item pai até a lista, para abas aninhadas
    linhas: list = field(default_factory=list)
    links: list = field(default_factory=list)   # (célula, aba destino, célula destino)


def oculto(c: dict) -> bool:
    v = c.get("visibility") or {}
    return v.get("type") == "FIXED" and v.get("value") is False


def coletavel(c: dict) -> bool:
    m = c.get("maximum")
    return m is None or (isinstance(m, (int, float)) and m > 1)


def aba_propria(c: dict, formato: str) -> bool:
    t = c.get("type")
    if formato == "csv":
        if t in SEMPRE_MULTI:
            return True
        if t in CSV_EM_LINHA:
            return False
        return coletavel(c)
    return t in SEMPRE_MULTI or (coletavel(c) and t not in NUNCA_MULTI)


def cod(rotulo: str) -> str:
    return f"Código ({rotulo})"


def colunas_folha(c: dict, anc: tuple) -> list[Col]:
    rot, t = c.get("label") or "", c.get("type")
    if t in MIDIA:
        return [Col("Nome original", "mid_nome", c, anc=anc), Col("Link para download", "mid_link", c, anc=anc),
                Col(rot, "mid_arq", c, anc=anc)]
    if t in COORD:
        return [Col(rot + " (Latitude)", "lat", c, anc=anc), Col(rot + " (Longitude)", "lon", c, anc=anc)]
    if t == "relational":
        cols = [Col(rot + " (Código)", "rel_cod", c, anc=anc)]
        if rel_tem_rotulo(c):
            cols.append(Col(rot + " (Rótulo)", "rel_rot", c, anc=anc))
        return cols
    return [Col(rot, "item", c, anc=anc)]


def colunas_simples(c: dict, prefixo: str, caminho: tuple, anc: tuple) -> list[Col]:
    rot, t = prefixo + (c.get("label") or ""), c.get("type")
    if t in COORD:
        return [Col(rot + " (Latitude)", "lat", c, caminho, anc=anc), Col(rot + " (Longitude)", "lon", c, caminho, anc=anc)]
    if t == "signature":
        return [Col(rot + " (Nome original)", "ass_nome", c, caminho, anc=anc),
                Col(rot + " (Link para download)", "ass_link", c, caminho, anc=anc),
                Col(rot + " (Nome do arquivo)", "ass_arq", c, caminho, anc=anc)]
    if t == "relational":
        cols = [Col(rot + " (Código)", "rel_cod", c, caminho, anc=anc)]
        if rel_tem_rotulo(c):
            cols.append(Col(rot + " (Rótulo)", "rel_rot", c, caminho, anc=anc))
        return cols
    return [Col(rot, "direto", c, caminho, anc=anc)]


@dataclass(eq=False)
class Esquema:
    estrutura: dict
    formato: str
    raiz: Aba
    abas: list                     # abas filhas em ordem (pré-ordem da estrutura)
    aba_unica: str | None = None
    csv_sep: str = ";"
    csv_decimal: str = ","
    rotulos: dict = field(default_factory=dict)
    fuso: ZoneInfo = field(default_factory=lambda: ZoneInfo("America/Sao_Paulo"))


def _andar_raiz(comps, caminho, prefixo, cab_cod, cols, abas, formato, anc):
    for c in comps or []:
        if c.get("type") == "separator" or oculto(c):
            continue
        a = anc + (c,)
        if c.get("type") == "group":
            if coletavel(c):
                aba = Aba(nome_aba(c.get("label")), c.get("label") or "", comp=c, grupo=True, caminho_topo=caminho)
                abas.append(aba)
                cols.append(Col(prefixo + (c.get("label") or ""), "qtd", c, caminho, filha=aba, anc=a))
                aba.cols = [Col(cab_cod, "cod_pai", anc=anc), Col(cod(c.get("label") or ""), "cod", anc=a)]
                aba.cols += _andar_grupo(c.get("components"), c.get("label") or "", (c.get("label") or "") + " > ",
                                         aba, (), abas, formato, a)
            else:
                _andar_raiz(c.get("components"), caminho + (c.get("id"),), prefixo + (c.get("label") or "") + " > ",
                            cod(c.get("label") or ""), cols, abas, formato, a)
            continue
        if aba_propria(c, formato):
            aba = Aba(nome_aba(c.get("label")), c.get("label") or "", comp=c, caminho_topo=caminho)
            aba.cols = [Col(cab_cod, "cod_pai", anc=anc), Col(cod(c.get("label") or ""), "cod", anc=a)] + colunas_folha(c, a)
            abas.append(aba)
            cols.append(Col(prefixo + (c.get("label") or ""), "qtd", c, caminho, filha=aba, anc=a))
            continue
        cols += colunas_simples(c, prefixo, caminho, a)


def _andar_grupo(comps, rotulo_pai, prefixo, aba_pai, caminho, abas, formato, anc) -> list[Col]:
    cols: list[Col] = []
    for c in comps or []:
        if c.get("type") == "separator" or oculto(c):
            continue
        a = anc + (c,)
        rot = c.get("label") or ""
        if c.get("type") == "group":
            if coletavel(c):
                naba = Aba(nome_aba(rot), rot, comp=c, grupo=True, pai=aba_pai, caminho_aninhado=caminho + (c.get("id"),))
                abas.append(naba)
                cols.append(Col(prefixo + rot, "qtd", c, caminho, filha=naba, anc=a))
                naba.cols = [Col(cod(rotulo_pai), "cod_pai", anc=anc), Col(cod(rot), "cod", anc=a)]
                naba.cols += _andar_grupo(c.get("components"), rot, rot + " > ", naba, (), abas, formato, a)
            else:
                cols += _andar_grupo(c.get("components"), rot, prefixo + rot + " > ", aba_pai,
                                     caminho + (c.get("id"),), abas, formato, a)
            continue
        if aba_propria(c, formato):
            naba = Aba(nome_aba(rot), rot, comp=c, pai=aba_pai, caminho_aninhado=caminho + (c.get("id"),))
            naba.cols = [Col(cod(rotulo_pai), "cod_pai", anc=anc), Col(cod(rot), "cod", anc=a)] + colunas_folha(c, a)
            abas.append(naba)
            cols.append(Col(prefixo + rot, "qtd", c, caminho, filha=naba, anc=a))
            continue
        cols += colunas_simples(c, prefixo, caminho, a)
    return cols


META = [("coleta_lat", "Local da coleta (Latitude)"), ("coleta_lon", "Local da coleta (Longitude)"),
        ("edicao_lat", "Local de edição (Latitude)"), ("edicao_lon", "Local de edição (Longitude)"),
        ("criado_por", "Criado por"), ("atualizado_por", "Atualizado por"), ("criado_em", "Criado em"),
        ("criado_em_disp", "Criado em (horário do dispositivo)"), ("atualizado_em", "Atualizado em")]


def _norm(s) -> str:
    return re.sub(r"\s+", " ", str(s or "")).strip().casefold()


def _caminho_rotulo(anc: tuple) -> str:
    return " > ".join((c.get("label") or "") for c in anc)


def _casa(entrada: str, col: Col) -> bool:
    """A entrada de mostrar/ocultar vale para o campo (rótulo, chave ou caminho 'Grupo > Campo') ou para
    um grupo acima dele; também vale o cabeçalho exato da coluna."""
    e = _norm(entrada)
    if e == _norm(col.rotulo):
        return True
    for i, c in enumerate(col.anc):
        if e in (_norm(c.get("label")), _norm(c.get("id")), _norm(_caminho_rotulo(col.anc[:i + 1]))):
            return True
    return False


def _todas_colunas(raiz: Aba, abas: list) -> list[Col]:
    return [c for a in [raiz, *abas] for c in a.cols]


def validar_ajustes(ajustes: dict | None) -> dict:
    aj = dict(ajustes or {})
    desconhecidas = sorted(set(aj) - CHAVES_AJUSTES)
    if desconhecidas:
        raise ErroAjuste(f"Ajuste desconhecido: {', '.join(desconhecidas)}. Válidos: {', '.join(sorted(CHAVES_AJUSTES))}.")
    campos = aj.get("campos") or {}
    if not isinstance(campos, dict) or set(campos) - {"mostrar", "ocultar"}:
        raise ErroAjuste("campos aceita só mostrar[] e ocultar[], com o rótulo ou a chave de cada campo.")
    for k in ("mostrar", "ocultar"):
        if campos.get(k) is not None and not isinstance(campos[k], list):
            raise ErroAjuste(f"campos.{k} precisa ser uma lista de rótulos ou chaves.")
    modo = aj.get("aba_unica")
    if modo is not None and modo not in MODOS_ABA_UNICA:
        raise ErroAjuste("aba_unica aceita 'linhas' (uma linha por resposta) ou 'colunas' (cada resposta numa coluna).")
    meta = aj.get("metadados", "fim")
    if meta not in POSICOES_META:
        raise ErroAjuste("metadados aceita 'fim' (padrão), 'inicio' ou 'fora'.")
    if not isinstance(aj.get("rotulos") or {}, dict):
        raise ErroAjuste("rotulos é um dicionário {nome atual: nome novo}.")
    csv = aj.get("csv") or {}
    if not isinstance(csv, dict) or set(csv) - {"separador", "decimal"}:
        raise ErroAjuste("csv aceita separador (';' ou ',') e decimal (',' ou '.').")
    sep = {"ponto_e_virgula": ";", "virgula": ","}.get(csv.get("separador"), csv.get("separador", ";"))
    dec = {"virgula": ",", "ponto": "."}.get(csv.get("decimal"), csv.get("decimal", ","))
    if sep not in (";", ",") or dec not in (",", "."):
        raise ErroAjuste("csv: separador ';' ou ',' e decimal ',' ou '.'.")
    extras = []
    for x in aj.get("extras") or []:
        k = SINONIMOS_EXTRAS.get(_norm(x), _norm(x))
        if k not in EXTRAS:
            raise ErroAjuste(f"Extra desconhecido: {x}. Válidos: precisao, altitude, origem.")
        if k not in extras:
            extras.append(k)
    return {"mostrar": campos.get("mostrar") or [], "ocultar": campos.get("ocultar") or [], "aba_unica": modo,
            "metadados": meta, "rotulos": aj.get("rotulos") or {}, "sep": sep, "dec": dec, "extras": extras}


def montar_esquema(estrutura: dict, formato: str, ajustes: dict | None = None,
                   fuso: ZoneInfo | None = None) -> Esquema:
    """Esquema da exportação (abas, colunas e cabeçalhos) pela estrutura; aplica os ajustes. Não lê a API."""
    aj = validar_ajustes(ajustes)
    nome = estrutura.get("name") or ""
    cab = cod(nome)
    cols: list[Col] = []
    abas: list[Aba] = []
    _andar_raiz(estrutura.get("components"), (), "", cab, cols, abas, formato, ())
    meta = []
    if estrutura.get("answer_tracking"):
        meta += [Col(r, "meta", meta=k) for k, r in META[:4]]
    meta += [Col(r, "meta", meta=k) for k, r in META[4:]]
    raiz = Aba(nome_aba(nome), nome, cols=[Col(cab, "cod"), *cols, *meta])
    nomes_unicos([raiz, *abas])
    esq = Esquema(estrutura, formato, raiz, abas, aj["aba_unica"], aj["sep"], aj["dec"], dict(aj["rotulos"]),
                  fuso or ZoneInfo("America/Sao_Paulo"))
    _aplicar(esq, aj)
    return esq


def _aplicar(esq: Esquema, aj: dict) -> None:
    todas = _todas_colunas(esq.raiz, esq.abas)
    for chave in ("mostrar", "ocultar"):
        faltam = [e for e in aj[chave] if not any(_casa(e, c) for c in todas if c.tipo not in ("cod", "cod_pai"))]
        if faltam:
            raise ErroAjuste(f"campos.{chave}: não achei {', '.join(map(str, faltam))} neste formulário "
                             "(use o rótulo como aparece na estrutura, a chave ou o cabeçalho da coluna).")
    rotulos_validos = {_norm(c.rotulo) for c in todas} | {_norm(x.get("label")) for c in todas for x in c.anc} \
        | {_norm(x.get("id")) for c in todas for x in c.anc} | {_norm(r) for _, r in META} | {_norm(r) for r in EXTRAS.values()}
    faltam = [k for k in aj["rotulos"] if _norm(k) not in rotulos_validos]
    if faltam:
        raise ErroAjuste(f"rotulos: não achei {', '.join(map(str, faltam))} (use o cabeçalho atual da coluna ou o rótulo do campo).")
    # chave técnica em rotulos vale como o rótulo do campo
    ids = {x.get("id"): x.get("label") for c in todas for x in c.anc}
    esq.rotulos = {(ids.get(k) or k): v for k, v in aj["rotulos"].items()}

    mostrar, ocultar = aj["mostrar"], aj["ocultar"]

    def fica(c: Col) -> bool:
        if c.tipo in ("cod", "cod_pai"):
            return True
        if c.tipo == "meta":
            return not any(_norm(e) == _norm(c.rotulo) for e in ocultar)
        if any(_casa(e, c) for e in ocultar):
            return False
        return not mostrar or any(_casa(e, c) for e in mostrar)

    def ordem(c: Col) -> int:
        if c.tipo == "qtd":
            return min([ordem(x) for x in c.filha.cols if x.tipo not in ("cod", "cod_pai")] or [len(mostrar)])
        for i, e in enumerate(mostrar):
            if _casa(e, c):
                return i
        return len(mostrar)

    vivas: set[int] = set()

    def filtrar(aba: Aba) -> bool:
        novas = []
        for c in aba.cols:
            if c.tipo == "qtd":
                # a aba filha fica se sobrou nela alguma coluna de valor (o campo pedido pode estar lá dentro)
                if not any(_casa(e, c) for e in ocultar) and filtrar(c.filha):
                    novas.append(c)
                continue
            if fica(c):
                novas.append(c)
        aba.cols = novas
        tem_valor = any(c.tipo not in ("cod", "cod_pai", "meta") for c in novas)
        if tem_valor or aba is esq.raiz:
            vivas.add(id(aba))
        return tem_valor

    filtrar(esq.raiz)
    esq.abas = [a for a in esq.abas if id(a) in vivas and (a.pai is None or id(a.pai) in vivas)]
    if mostrar:
        for aba in [esq.raiz, *esq.abas]:
            fixas = [c for c in aba.cols if c.tipo in ("cod", "cod_pai")]
            meio = sorted([c for c in aba.cols if c.tipo not in ("cod", "cod_pai", "meta")], key=ordem)
            aba.cols = fixas + meio + [c for c in aba.cols if c.tipo == "meta"]
    # metadados e extras
    r = esq.raiz
    meta = [c for c in r.cols if c.tipo == "meta"]
    corpo = [c for c in r.cols if c.tipo != "meta"]
    if aj["metadados"] == "fora":
        meta = []
    meta += [Col(EXTRAS[k], "meta", meta="x_" + k) for k in aj["extras"]]
    r.cols = [corpo[0], *meta, *corpo[1:]] if aj["metadados"] == "inicio" else [*corpo, *meta]


# --------------------------------------------------------------------------------------------
# Preenchimento
# --------------------------------------------------------------------------------------------

def _resolver(escopo, caminho: tuple, chave):
    m = escopo
    for g in caminho:
        m = m.get(g) if isinstance(m, dict) else None
        if not isinstance(m, dict):
            return None
    return m.get(chave) if isinstance(m, dict) else None


def _seguir(item, caminho: tuple) -> list:
    m = item
    for g in caminho[:-1]:
        m = m.get(g) if isinstance(m, dict) else None
    v = m.get(caminho[-1]) if isinstance(m, dict) and caminho else None
    return v if isinstance(v, list) else []


def _celula(col: Col, escopo, fuso: ZoneInfo) -> Cel | None:
    c = col.comp or {}
    if col.tipo in ("item", "mid_nome", "mid_link", "mid_arq"):
        v = escopo
    else:
        v = _resolver(escopo, col.caminho, c.get("id"))
    t = col.tipo
    if t == "qtd":
        return Cel("num", float(len(v)) if isinstance(v, list) else 0.0, None, "int")
    if t in ("lat", "lon"):
        xy = coord(v)
        return Cel("num", arredonda_coord(xy[0] if t == "lat" else xy[1])) if xy else VAZIA
    if t in ("ass_nome", "ass_link", "ass_arq"):
        if not isinstance(v, str) or not v:
            return VAZIA
        original, bonito = partes_midia(v)
        return texto({"ass_nome": original, "ass_link": v, "ass_arq": bonito}[t])
    if t in ("mid_nome", "mid_link", "mid_arq"):
        url = go_v(v)
        original, bonito = partes_midia(url)
        return texto({"mid_nome": original, "mid_link": url, "mid_arq": bonito}[t])
    if t in ("rel_cod", "rel_rot"):
        r = relacional(v)
        return texto(r[0] if t == "rel_cod" else r[1]) if r else VAZIA
    return celula_folha(c.get("type") or "", v, fuso)


def _meta(col: Col, p: dict, fuso: ZoneInfo) -> Cel | None:
    m = p.get("meta_data") or {}
    k = col.meta
    if k in ("coleta_lat", "coleta_lon", "edicao_lat", "edicao_lon"):
        xy = coord(m.get("created_at_coordinates" if k.startswith("coleta") else "updated_at_coordinates"))
        return Cel("num", arredonda_coord(xy[0] if k.endswith("lat") else xy[1])) if xy else VAZIA
    if k == "criado_por":
        return texto(m.get("created_by_user_name") if m.get("created_by_user_name") is not None else "Anônimo")
    if k == "atualizado_por":
        return texto(m.get("updated_by_user_name")) if m.get("updated_by_user_name") is not None else VAZIA
    if k in ("criado_em", "criado_em_disp"):
        d = ler_data(m.get("created_at" if k == "criado_em" else "created_at_device"))
        return Cel("datahora", d.astimezone(fuso)) if d else VAZIA
    if k == "atualizado_em":
        d, criado = ler_data(m.get("updated_at")), ler_data(m.get("created_at"))
        return Cel("datahora", d.astimezone(fuso)) if d and d != criado else VAZIA
    props = (m.get("created_at_coordinates") or {}).get("properties") or {} \
        if isinstance(m.get("created_at_coordinates"), dict) else {}
    if k == "x_precisao":
        return Cel("num", float(props["precision"])) if num(props.get("precision")) is not None else VAZIA
    if k == "x_altitude":
        return Cel("num", float(props["elevation"])) if num(props.get("elevation")) is not None else VAZIA
    if k == "x_origem":
        s = m.get("created_at_source")
        return texto(FONTES.get(s, s)) if s else VAZIA
    return VAZIA


def _letra(n: int) -> str:
    s = ""
    while n > 0:
        n, r = divmod(n - 1, 26)
        s = chr(65 + r) + s
    return s


def preencher(esq: Esquema, preenchimentos: list[dict]) -> None:
    """Uma linha por preenchimento na raiz e uma por item nas abas filhas, com os códigos e os links."""
    filhas_de: dict[int, list[Aba]] = {}
    for a in esq.abas:
        filhas_de.setdefault(id(a.pai) if a.pai else 0, []).append(a)
    raiz = esq.raiz
    col_qtd = {id(c.filha): i for i, c in enumerate(raiz.cols) if c.tipo == "qtd"}

    def linha_item(aba: Aba, item, cod_pai: str, cod_proprio: str) -> list:
        linha = [texto(cod_pai), texto(cod_proprio)]
        for c in aba.cols[2:]:
            if aba.grupo:
                linha.append(_celula(c, item, esq.fuso) if isinstance(item, dict) else VAZIA)
            else:
                linha.append(_celula_item(c, item))
        return linha

    def _celula_item(c: Col, item):
        if c.tipo in ("lat", "lon"):
            xy = coord(item)
            return Cel("num", arredonda_coord(xy[0] if c.tipo == "lat" else xy[1])) if xy else VAZIA
        if c.tipo in ("rel_cod", "rel_rot"):
            r = relacional(item)
            return texto(r[0] if c.tipo == "rel_cod" else r[1]) if r else VAZIA
        if c.tipo == "item":
            return celula_folha((c.comp or {}).get("type") or "", item, esq.fuso)
        return _celula(c, item, esq.fuso)

    def encher(aba: Aba, itens: list, cod_pai: str, aba_pai: str, linha_pai: int) -> None:
        for i, item in enumerate(itens):
            proprio = f"{cod_pai}.{i}"
            aba.linhas.append(linha_item(aba, item, cod_pai, proprio))
            n = len(aba.linhas) + 1
            aba.links.append((f"A{n}", aba_pai, f"A{linha_pai}"))
            if isinstance(item, dict):
                for f in filhas_de.get(id(aba), []):
                    encher(f, _seguir(item, f.caminho_aninhado), proprio, aba.nome, n)

    for ip, p in enumerate(preenchimentos):
        n_raiz = ip + 2
        resp = p.get("answer") if isinstance(p.get("answer"), dict) else {}
        pid = str(p.get("id") or "")
        linha = []
        for c in raiz.cols:
            if c.tipo == "cod":
                linha.append(texto(pid))
            elif c.tipo == "meta":
                linha.append(_meta(c, p, esq.fuso))
            else:
                linha.append(_celula(c, resp, esq.fuso))
        raiz.linhas.append(linha)
        for a in filhas_de.get(0, []):
            itens = _resolver(resp, a.caminho_topo, (a.comp or {}).get("id"))
            if not isinstance(itens, list) or not itens:
                continue
            if id(a) in col_qtd:
                raiz.links.append((f"{_letra(col_qtd[id(a)] + 1)}{n_raiz}", a.nome, f"A{len(a.linhas) + 2}"))
            encher(a, itens, pid, raiz.nome, n_raiz)


# --------------------------------------------------------------------------------------------
# Aba única: modo linhas e modo colunas
# --------------------------------------------------------------------------------------------

SUFIXO_FOLHA = {"mid_nome": " (Nome original)", "mid_link": " (Link para download)", "mid_arq": " (Nome do arquivo)",
                "lat": " (Latitude)", "lon": " (Longitude)", "rel_cod": " (Código)", "rel_rot": " (Rótulo)", "item": ""}


def _cab_filho(filha: Aba, col: Col, base: str) -> str:
    """Cabeçalho de uma coluna da aba filha, pendurado no cabeçalho da coluna de contagem (base)."""
    if filha.grupo:
        prefixo = (filha.comp.get("label") or "") + " > "
        resto = col.rotulo[len(prefixo):] if col.rotulo.startswith(prefixo) else col.rotulo
        return f"{base} > {resto}"
    return base + SUFIXO_FOLHA.get(col.tipo, "")


def _indices(esq: Esquema) -> dict[int, dict[str, list]]:
    idx: dict[int, dict[str, list]] = {}
    for a in esq.abas:
        d: dict[str, list] = {}
        for linha in a.linhas:
            d.setdefault(linha[0].valor if linha[0] else "", []).append(linha)
        idx[id(a)] = d
    return idx


def aviso_aba_unica(esq: Esquema) -> str | None:
    """Aba única foi pensada para formulário simples: campos simples, poucas múltiplas escolhas ou um grupo
    repetível sem subgrupo. Fora disso, funciona pela regra simples, mas o padrão com abas lê melhor."""
    if not esq.aba_unica:
        return None
    topo = [a for a in esq.abas if a.pai is None and a.grupo]
    if any(a.pai is not None for a in esq.abas) or len(topo) > 1:
        return ("Este formulário tem respostas múltiplas dentro de grupo repetível, ou mais de um grupo repetível: numa aba só as respostas "
                + ("ficam lado a lado por posição" if esq.aba_unica == "linhas" else "viram muitas colunas")
                + ". O padrão com abas representa melhor.")
    return None


def aba_unica(esq: Esquema) -> tuple[list[tuple[str, Col]], list[list]]:
    """Achata raiz e abas filhas numa aba só. Devolve (cabeçalhos com a coluna de origem, linhas)."""
    idx = _indices(esq)
    raiz = esq.raiz

    def proprio(aba: Aba, linha: list) -> str:
        c = linha[0] if aba is raiz else linha[1]
        return c.valor if c else ""

    def valor_cols(aba: Aba) -> list[tuple[int, Col]]:
        inicio = 1 if aba is raiz else 2
        return list(enumerate(aba.cols))[inicio:]

    if esq.aba_unica == "colunas":
        maximo: dict[int, int] = {id(a): max([len(v) for v in idx[id(a)].values()] or [0]) for a in esq.abas}

        def layout(aba: Aba, base: str | None) -> list:
            saida = []
            for i, c in valor_cols(aba):
                cab = c.rotulo if base is None else _cab_filho(aba, c, base)
                if c.tipo == "qtd":
                    for n in range(1, maximo[id(c.filha)] + 1):
                        saida.append(("filha", i, c.filha, n, layout(c.filha, f"{cab} {n}")))
                else:
                    saida.append(("col", i, c, cab))
            return saida

        plano = layout(raiz, None)

        def cabecalhos(pl) -> list:
            out = []
            for e in pl:
                out += [(e[3], e[2])] if e[0] == "col" else cabecalhos(e[4])
            return out

        def valores(pl, aba: Aba, linha: list | None) -> list:
            out = []
            for e in pl:
                if e[0] == "col":
                    out.append(linha[e[1]] if linha is not None else VAZIA)
                else:
                    filhos = idx[id(e[2])].get(proprio(aba, linha), []) if linha is not None else []
                    out += valores(e[4], e[2], filhos[e[3] - 1] if e[3] <= len(filhos) else None)
            return out

        cabs = [(raiz.cols[0].rotulo, raiz.cols[0])] + cabecalhos(plano)
        linhas = [[l[0]] + valores(plano, raiz, l) for l in raiz.linhas]
        return cabs, linhas

    # modo linhas: itens de grupo repetível um embaixo do outro; campos multivalorados do mesmo nível lado a lado
    def layout_l(aba: Aba, base: str | None) -> list:
        saida = []
        for i, c in valor_cols(aba):
            cab = c.rotulo if base is None else _cab_filho(aba, c, base)
            if c.tipo == "qtd":
                saida.append(("filha", i, c.filha, layout_l(c.filha, cab)))
            else:
                saida.append(("col", i, c, cab))
        return saida

    plano = layout_l(raiz, None)

    def cabecalhos_l(pl) -> list:
        out = []
        for e in pl:
            out += [(e[3], e[2])] if e[0] == "col" else cabecalhos_l(e[3])
        return out

    def largura(pl) -> int:
        return sum(1 if e[0] == "col" else largura(e[3]) for e in pl)

    def linhas_de(pl, aba: Aba, linha: list) -> list[list]:
        base, blocos = [], []
        for e in pl:
            if e[0] == "col":
                base.append(("v", linha[e[1]]))
            else:
                filhos = idx[id(e[2])].get(proprio(aba, linha), [])
                sub = [r for f in filhos for r in linhas_de(e[3], e[2], f)]
                base.append(("b", len(blocos), largura(e[3])))
                blocos.append(sub)
        n = max([1] + [len(b) for b in blocos])
        saida = []
        for k in range(n):
            r = []
            for item in base:
                if item[0] == "v":
                    r.append(item[1])
                else:
                    b = blocos[item[1]]
                    r += b[k] if k < len(b) else [VAZIA] * item[2]
            saida.append(r)
        return saida

    cabs = [(raiz.cols[0].rotulo, raiz.cols[0])] + cabecalhos_l(plano)
    linhas = [[l[0]] + r for l in raiz.linhas for r in linhas_de(plano, raiz, l)]
    return cabs, linhas


# --------------------------------------------------------------------------------------------
# Rótulos novos (ajuste rotulos)
# --------------------------------------------------------------------------------------------

def renomear(cab: str, col: Col | None, rotulos: dict) -> str:
    if not rotulos:
        return cab
    exato = {_norm(k): v for k, v in rotulos.items()}
    if _norm(cab) in exato:
        return exato[_norm(cab)]
    labels = {(x.get("label") or "") for x in (col.anc if col else ())}
    trocas = [(k, v) for k, v in rotulos.items() if k in labels]
    if not trocas:
        return cab
    partes = cab.split(" > ")
    for k, v in trocas:
        pad = re.compile(rf"^(Código \()?{re.escape(k)}(?=$|\)| \(| \d)")
        partes = [pad.sub(lambda m: (m.group(1) or "") + v, p) for p in partes]
    return " > ".join(partes)


# --------------------------------------------------------------------------------------------
# Gravação
# --------------------------------------------------------------------------------------------

def _serial(d: datetime) -> float:
    """Número de série do Excel pelo relógio de parede, como toXLSXSerial do serviço."""
    ingenuo = datetime(d.year, d.month, d.day, d.hour, d.minute, d.second, d.microsecond)
    delta = ingenuo - datetime(1899, 12, 30)
    ns = (delta.days * 86400 + delta.seconds) * 10**9 + delta.microseconds * 1000
    horas = ns // (3600 * 10**9)
    resto = ns % (3600 * 10**9)
    return (float(horas) + resto / (60 * 60 * 1e9)) / 24.0


def _largura(cab: str, celulas: list) -> float:
    w = float(len(cab.encode("utf-8")))
    for c in celulas:
        if c is None:
            continue
        l = {"datahora": 16, "data": 10, "num": 12}.get(c.tipo, len(str(c.valor).encode("utf-8")))
        w = max(w, l)
    return min(60.0, max(8.0, w + 2.0))


def _estilo_cabecalho():
    from openpyxl.styles import Alignment, Font, PatternFill
    return (Font(b=True, color=BRANCO, sz=11), PatternFill("solid", fgColor=VERDE),
            Alignment(horizontal="center", vertical="center"))


def _escrever_aba(ws, cabs: list[str], linhas: list[list], links: list, formatos: tuple[str, str]) -> None:
    from openpyxl.worksheet.hyperlink import Hyperlink
    fonte, fundo, alinha = _estilo_cabecalho()
    f_data, f_dh = formatos
    for j, h in enumerate(cabs, 1):
        cel = ws.cell(row=1, column=j, value=h)
        cel.font, cel.fill, cel.alignment = fonte, fundo, alinha
    for i, linha in enumerate(linhas, 2):
        for j, c in enumerate(linha, 1):
            if c is None:
                continue
            cel = ws.cell(row=i, column=j)
            if c.tipo in ("data", "datahora"):
                cel.value = _serial(c.valor)
                cel.number_format = f_data if c.tipo == "data" else f_dh
            elif c.tipo == "num":
                cel.value = c.valor
                if c.csv == "moeda":
                    s = simbolo(c.moeda)
                    cel.number_format = (f'"{s}"' if s else "") + "#,##0.00"
                elif c.xfmt:
                    cel.number_format = c.xfmt
            else:
                cel.value = c.valor
                cel.data_type = "s"
    for ref, aba, alvo in links:
        ws[ref].hyperlink = Hyperlink(ref=ref, location="'" + aba.replace("'", "''") + "'!" + alvo)
    for j, h in enumerate(cabs, 1):
        ws.column_dimensions[_letra(j)].width = _largura(h, [l[j - 1] if j - 1 < len(l) else None for l in linhas])


def _linhas_leia_me(filtros: list[str], avisos: list[str]) -> int:
    n = 0
    for bloco in (filtros, avisos):
        if bloco:
            n += (1 if n else 0) + 1 + len(bloco)
    return n


def _leia_me(ws, nome_form: str, info: str, blocos: list[tuple[str, list[str]]]) -> None:
    from copy import copy

    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.styles.fonts import DEFAULT_FONT
    verde = PatternFill("solid", fgColor=VERDE)
    centro = Alignment(horizontal="center", vertical="center", wrap_text=True)
    normal = Font(b=True, color=BRANCO, sz=11)
    ws.column_dimensions["A"].width = 3
    for col in "BCDEFGHI":
        ws.column_dimensions[col].width = 11
    ws.column_dimensions["J"].width = 3
    for r in range(7, 16):
        ws.row_dimensions[r].height = 20
    ultima = max(TOPO_LEIA_ME, 24 + _linhas_leia_me(*(b[1] for b in blocos)))
    for r in range(1, ultima + 1):
        for c in range(1, 12):
            cel = ws.cell(row=r, column=c)
            cel.font, cel.fill, cel.alignment = normal, verde, centro
    ws.merge_cells("B7:I15")
    ws["B7"].fill = PatternFill("solid", fgColor=BRANCO)
    ws["B7"].font = copy(DEFAULT_FONT)
    ws["B7"].alignment = Alignment(horizontal="center", vertical="center")
    ws.merge_cells("B17:I19")
    ws["B17"] = nome_form
    ws["B17"].font = Font(b=True, color=BRANCO, sz=14)
    ws.merge_cells("B21:I23")
    ws["B21"] = info
    linha = 25
    esquerda = Alignment(horizontal="left", vertical="center", wrap_text=True, indent=1)
    for titulo, itens in blocos:
        if not itens:
            continue
        if linha > 25:
            linha += 1
        ws.merge_cells(f"B{linha}:I{linha}")
        ws[f"B{linha}"] = titulo
        for k, item in enumerate(itens, 1):
            ws.merge_cells(f"B{linha + k}:I{linha + k}")
            cel = ws[f"B{linha + k}"]
            cel.value = item
            cel.font, cel.alignment = Font(color=BRANCO, sz=11), esquerda
        linha += 1 + len(itens)


@dataclass
class Contexto:
    agora: datetime
    gerado_por: str
    filtros: list[str]
    aviso_incompleto: str | None = None


def _info_leia_me(esq: Esquema, ctx: Contexto, total: int) -> str:
    return (f"Exportação realizada por {ctx.gerado_por} em {ctx.agora.strftime('%d/%m/%Y')} às "
            f"{ctx.agora.strftime('%H:%M')} · Total de respostas: {total}")


def _abas_para_gravar(esq: Esquema) -> list[tuple[str, list[str], list[list], list]]:
    """(nome da aba ou base do arquivo, cabeçalhos, linhas, links) no formato pedido."""
    r = esq.rotulos
    if esq.aba_unica:
        cabs, linhas = aba_unica(esq)
        nome = esq.raiz.nome if esq.formato == "xlsx" else esq.raiz.rotulo
        return [(nome, [renomear(h, c, r) for h, c in cabs], linhas, [])]
    saida = []
    for a in [esq.raiz, *esq.abas]:
        idx = [i for i, c in enumerate(a.cols) if not (esq.formato == "csv" and c.tipo == "qtd")]
        cabs = [renomear(a.cols[i].rotulo, a.cols[i], r) for i in idx]
        linhas = [[l[i] for i in idx] for l in a.linhas]
        saida.append((a.nome if esq.formato == "xlsx" else a.rotulo, cabs, linhas, a.links if esq.formato == "xlsx" else []))
    return saida


def gravar_xlsx(esq: Esquema, destino: Path, ctx: Contexto) -> tuple[Path, dict]:
    from openpyxl import Workbook

    e = esq.estrutura
    arquivo = destino / (nome_arquivo(e.get("name") or "", str(e.get("version") or ""), ctx.agora) + ".xlsx")
    wb = Workbook()
    leia = wb.active
    leia.title = "LEIA-ME"
    blocos = [("Filtros utilizados:", ctx.filtros),
              ("Exportação incompleta:", [ctx.aviso_incompleto] if ctx.aviso_incompleto else [])]
    _leia_me(leia, e.get("name") or "", _info_leia_me(esq, ctx, len(esq.raiz.linhas)), blocos)
    formatos = ("dd/mm/yyyy", "dd/mm/yyyy hh:mm")
    contagens = {}
    for nome, cabs, linhas, links in _abas_para_gravar(esq):
        ws = wb.create_sheet(nome)
        _escrever_aba(ws, cabs, linhas, links, formatos)
        contagens[nome] = len(linhas)
    leia.sheet_view.tabSelected = False
    wb.active = 1
    wb.worksheets[1].sheet_view.tabSelected = True
    destino.mkdir(parents=True, exist_ok=True)
    wb.save(arquivo)
    return arquivo, contagens


def _csv_texto(c: Cel | None, dec: str, fuso: ZoneInfo) -> str:
    if c is None:
        return ""
    if c.tipo == "data":
        return c.valor.strftime("%d/%m/%Y")
    if c.tipo == "datahora":
        return c.valor.strftime("%d/%m/%Y %H:%M")
    if c.tipo == "num":
        if c.csv == "int":
            return go_f(c.valor)
        if c.csv == "moeda":
            s = f"{c.valor:.2f}"
            s = s.replace(".", ",", 1) if dec == "," else s
            sim = simbolo(c.moeda)
            return (sim or "") + s
        s = go_f(c.valor)
        return s.replace(".", ",", 1) if dec == "," else s
    return str(c.valor)


def _csv_bytes(cabs: list[str], linhas: list[list], sep: str, dec: str, fuso: ZoneInfo) -> bytes:
    def reg(campos):
        return sep.join('"' + str(x).replace('"', '""') + '"' for x in campos) + "\n"
    corpo = reg(cabs) + "".join(reg(_csv_texto(c, dec, fuso) for c in l) for l in linhas)
    return b"\xef\xbb\xbf" + corpo.encode("utf-8")


def leia_me_txt(esq: Esquema, ctx: Contexto) -> str:
    t = "Exportação Coletum"
    linhas = [t, "=" * len(t), "", f"Formulário: {esq.estrutura.get('name') or ''}"]
    if ctx.gerado_por:
        linhas.append(f"Gerado por: {ctx.gerado_por}")
    linhas += [f"Data de exportação: {ctx.agora.strftime('%d/%m/%Y %H:%M')}",
               f"Total de preenchimentos: {len(esq.raiz.linhas)}"]
    texto_ = "\n".join(linhas) + "\n"
    if ctx.filtros:
        texto_ += "\nFiltros utilizados:\n" + "".join(f"  {f}\n" for f in ctx.filtros)
    if ctx.aviso_incompleto:
        texto_ += "\nExportação incompleta:\n" + f"  {ctx.aviso_incompleto}\n"
    return texto_


def gravar_csv(esq: Esquema, destino: Path, ctx: Contexto) -> tuple[Path, dict]:
    e = esq.estrutura
    versao = str(e.get("version") or "")
    pasta = destino / nome_arquivo(e.get("name") or "", versao, ctx.agora)
    pasta.mkdir(parents=True, exist_ok=True)
    usados: set[str] = set()
    contagens = {}
    for rotulo, cabs, linhas, _ in _abas_para_gravar(esq):
        base = nome_arquivo(rotulo, versao, ctx.agora)
        nome, n = base, 1
        while nome in usados:
            nome = f"{base}-{n:02d}"
            n += 1
        usados.add(nome)
        (pasta / f"{nome}.csv").write_bytes(_csv_bytes(cabs, linhas, esq.csv_sep, esq.csv_decimal, esq.fuso))
        contagens[f"{nome}.csv"] = len(linhas)
    (pasta / "LEIA-ME.txt").write_text(leia_me_txt(esq, ctx), encoding="utf-8", newline="\n")
    return pasta, contagens


def linhas_de_filtro(params: dict) -> list[str]:
    """Filtros da chamada no formato do LEIA-ME da exportação ('Campo | Operador | Valor')."""
    def dia(s: str) -> str:
        try:
            return datetime.strptime(str(s)[:10], "%Y-%m-%d").strftime("%d/%m/%Y")
        except ValueError:
            return str(s)
    saida = []
    for chave, rot, op in (("created_after", "Criado em", "Maior que"), ("created_before", "Criado em", "Menor que"),
                           ("updated_after", "Atualizado em", "Maior que"), ("updated_before", "Atualizado em", "Menor que")):
        if params.get(chave):
            saida.append(f"{rot} | {op} | {dia(params[chave])}")
    if params.get("created_at_source"):
        saida.append(f"Plataforma | Igual a | {FONTES.get(params['created_at_source'], params['created_at_source'])}")
    if params.get("created_by"):
        saida.append(f"Criado por | Igual a | {params['created_by']}")
    if params.get("updated_by"):
        saida.append(f"Atualizado por | Igual a | {params['updated_by']}")
    return saida


def exportar(estrutura: dict, preenchimentos: list[dict], formato: str, destino: Path, ctx: Contexto,
             ajustes: dict | None = None, fuso: ZoneInfo | None = None, esquema: Esquema | None = None) -> tuple[Path, dict]:
    esq = esquema or montar_esquema(estrutura, formato, ajustes, fuso)
    preencher(esq, preenchimentos)
    return gravar_xlsx(esq, destino, ctx) if formato == "xlsx" else gravar_csv(esq, destino, ctx)
